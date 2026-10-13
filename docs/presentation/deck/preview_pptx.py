"""Approximate check of the hybrid PPTX: redraw each slide from the saved file (background picture +
text boxes with their saved font, size, weight, colour, spacing, line spacing) as HTML, so the
pages can be rendered and compared with the original design. Not PowerPoint's own renderer."""
import base64
import html
import sys
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
PX = 6350
src = HERE.parents[2] / "docs" / "submission" / "assets" / "aura-health-deck.pptx"
out = Path(sys.argv[1]); out.mkdir(parents=True, exist_ok=True)
font = (HERE.parents[2] / "aura_ui" / "media" / "nunito.woff2").resolve()
deck = Presentation(src)
for number, slide in enumerate(deck.slides, start=1):
    parts = [f"<!doctype html><html><head><meta charset='utf-8'><style>@font-face{{font-family:Nunito;src:url('file://{font}') format('woff2');font-weight:200 1000}}"
             "body{margin:0;width:1920px;height:1080px;position:relative;overflow:hidden}.t{position:absolute;font-family:Nunito}</style></head><body>"]
    for shape in slide.shapes:
        if shape.shape_type == 13:  # picture
            data = base64.b64encode(shape.image.blob).decode()
            parts.append(f"<img src='data:image/jpeg;base64,{data}' style='position:absolute;left:0;top:0;width:1920px;height:1080px'>")
            continue
        if not shape.has_text_frame:
            continue
        x, y, w = shape.left / PX, shape.top / PX, shape.width / PX
        paragraphs = []
        for paragraph in shape.text_frame.paragraphs:
            align = {2: "center", 3: "right"}.get(int(paragraph.alignment) if paragraph.alignment is not None else 1, "left")
            line = paragraph.line_spacing.pt * 2 if paragraph.line_spacing is not None else None
            runs = []
            for run in paragraph.runs:
                f = run.font
                spc = run._r.get_or_add_rPr().get("spc")
                letter = f"letter-spacing:{int(spc) / 100 * 2}px;" if spc else ""
                color = f"#{f.color.rgb}" if f.color and f.color.type is not None else "#000"
                runs.append(f"<span style='font-size:{f.size.pt * 2}px;font-weight:{800 if f.bold else 400};color:{color};{letter}'>{html.escape(run.text)}</span>")
            paragraphs.append(f"<div style='text-align:{align};{f'line-height:{line}px;' if line else ''}'>{''.join(runs) or '&nbsp;'}</div>")
        parts.append(f"<div class='t' style='left:{x}px;top:{y}px;width:{w}px'>{''.join(paragraphs)}</div>")
    parts.append("</body></html>")
    (out / f"preview-{number:02d}.html").write_text("".join(parts))
print("ok")
