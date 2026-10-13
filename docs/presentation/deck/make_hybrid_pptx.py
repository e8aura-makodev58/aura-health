"""Hybrid deck, step 2: picture background per slide + every text block as an editable text box.

Positions, sizes, weights, colours, letter spacing, line height and alignment come from the
rendered slide pages (hybrid/text.json from hybrid.cjs). 1920 px = 13.333 in = 960 pt, so
1 px = 0.5 pt = 6350 EMU. Text uses the app's font, Nunito: install it (free, Google Fonts) for
the intended look; PowerPoint substitutes a similar font otherwise.
"""
import io
import json
import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE.parents[2] / "docs" / "submission" / "assets" / "aura-health-deck.pptx"
PX = 6350  # EMU per CSS pixel on a 1920-px-wide slide
FONT = "Nunito"
ALIGN = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}


def rgb(css):
    numbers = [int(float(n)) for n in re.findall(r"[\d.]+", css)[:3]]
    return RGBColor(*numbers)


def set_spacing(run, pixels):
    if pixels:
        run._r.get_or_add_rPr().set("spc", str(int(round(pixels * 0.5 * 100))))  # hundredths of a point


def background(slide, png):
    image = Image.open(png).convert("RGB")
    buffer = io.BytesIO(); image.save(buffer, "JPEG", quality=88, optimize=True); buffer.seek(0)
    picture = slide.shapes.add_picture(buffer, 0, 0, width=Emu(1920 * PX), height=Emu(1080 * PX))
    picture.name = "Background (art, cards, screenshots)"


def text_box(slide, block):
    # A little extra width absorbs font-metric differences so lines do not re-wrap.
    width = block["w"] * 1.04 + 4
    if block["align"] in ("center",):
        left = block["x"] - (width - block["w"]) / 2
    elif block["align"] in ("right", "end"):
        left = block["x"] - (width - block["w"])
    else:
        left = block["x"]
    shape = slide.shapes.add_textbox(Emu(int(left * PX)), Emu(int(block["y"] * PX)),
                                     Emu(int(width * PX)), Emu(int(max(block["h"], block["lineHeight"]) * PX)))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, side, 0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = ALIGN.get(block["align"], PP_ALIGN.LEFT)
    paragraph.line_spacing = Pt(block["lineHeight"] * 0.5)
    for item in block["runs"]:
        if item.get("br"):
            paragraph = frame.add_paragraph()
            paragraph.alignment = ALIGN.get(block["align"], PP_ALIGN.LEFT)
            paragraph.line_spacing = Pt(block["lineHeight"] * 0.5)
            continue
        run = paragraph.add_run()
        run.text = item["text"]
        font = run.font
        font.name = FONT
        font.size = Pt(round(item["size"] * 0.5, 1))
        font.bold = item["weight"] >= 600
        font.color.rgb = rgb(item["color"])
        set_spacing(run, item["spacing"])
    shape.name = "Text: " + "".join(r.get("text", " ") for r in block["runs"])[:40]


def main():
    slides = json.loads((HERE / "hybrid" / "text.json").read_text())
    notes = {row["slide"]: row["notes"] for row in json.loads((HERE / "slides.json").read_text())}
    deck = Presentation()
    deck.slide_width, deck.slide_height = Emu(1920 * PX), Emu(1080 * PX)
    blank = deck.slide_layouts[6]
    for item in slides:
        slide = deck.slides.add_slide(blank)
        background(slide, HERE / "hybrid" / f"bg-{item['slide']:02d}.png")
        for block in item["blocks"]:
            text_box(slide, block)
        slide.notes_slide.notes_text_frame.text = notes.get(item["slide"], "")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(OUT)
    print(OUT, OUT.stat().st_size, "bytes,", len(slides), "slides")


if __name__ == "__main__":
    main()
