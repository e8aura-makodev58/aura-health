"""Assemble png/slide-NN.png into a 16:9 PPTX with the scene narration as speaker notes.

Each slide is the full-bleed rendered image (JPEG, quality 90, to stay well under the form's
10 MB limit). Needs python-pptx (which brings Pillow); see README.md.
"""

import io
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
OUT = HERE.parent.parent / "submission" / "assets" / "aura-health-deck.pptx"


def main():
    manifest = json.loads((HERE / "slides.json").read_text())
    deck = Presentation()
    deck.slide_width, deck.slide_height = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in, 16:9
    blank = deck.slide_layouts[6]
    for row in manifest:
        image = Image.open(HERE / "png" / f"slide-{row['slide']:02d}.png").convert("RGB")
        buffer = io.BytesIO()
        image.save(buffer, "JPEG", quality=90, optimize=True)
        buffer.seek(0)
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(buffer, 0, 0, width=deck.slide_width, height=deck.slide_height)
        slide.notes_slide.notes_text_frame.text = row["notes"]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(OUT)
    print(f"{OUT} · {len(manifest)} slides · {OUT.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
