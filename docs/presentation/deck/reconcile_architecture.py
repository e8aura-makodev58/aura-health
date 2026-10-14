"""Refresh native diagrams in the existing deck and export the architecture PDF.

Requires python-pptx, reportlab, fonttools, brotli and pymupdf in a tooling environment.
Uses native_diagrams.py for both exports. PDF shapes/text are drawn with ReportLab;
this is not a PowerPoint/LibreOffice render. The reference PDF supplies only its
first/last pages, preserving the user's cover and close. No app dependencies.
"""
import argparse
from pathlib import Path
from tempfile import TemporaryDirectory

import pymupdf
from fontTools.ttLib import TTFont as SourceFont
from fontTools.varLib.instancer import instantiateVariableFont
from pptx import Presentation
from pptx.oxml.ns import qn
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

from native_diagrams import DIAGRAMS, PX

ROOT = Path(__file__).resolve().parents[3]
ASSETS = ROOT / 'docs/submission/assets'
KINDS = ('architecture', 'agent-loop', 'safety-precedence')
IMAGES = ('system-overview', 'agent-investigation', 'safety-precedence')


def site(shape, index):
    x, y, w, h = (value / PX / 2 for value in
                  (shape.left, shape.top, shape.width, shape.height))
    return ((x+w/2, y), (x, y+h/2), (x+w/2, y+h), (x+w, y+h/2))[index]


def export_diagrams(slides, output):
    """Render native rounded boxes, text and attached orthogonal connectors."""
    with TemporaryDirectory() as temp:
        for weight, name in ((400, 'Nunito'), (700, 'NunitoBold')):
            font = SourceFont(ROOT / 'aura_ui/media/nunito.woff2')
            font = instantiateVariableFont(font, {'wght': weight}, inplace=True)
            font.flavor = None
            # Distinct PostScript names avoid sharing the regular font's PDF subset.
            for record in font['name'].names:
                if record.nameID in (1, 4, 6):
                    record.string = name.encode(record.getEncoding())
            target = Path(temp) / (name + '.ttf')
            font.save(target)
            pdfmetrics.registerFont(TTFont(name, str(target)))
        canvas = Canvas(str(output), pagesize=(960, 540), pageCompression=1)
        canvas.setTitle('AURA Health — architecture reviewed against 90e31ea')
        for slide in slides:
            canvas.setFillColor(HexColor('#fbfaf7'))
            canvas.rect(0, 0, 960, 540, fill=1, stroke=0)
            shapes = {shape.shape_id: shape for shape in slide.shapes}
            for shape in slide.shapes:
                connector = shape._element.find(qn('p:nvCxnSpPr'))
                if connector is not None:
                    ends = connector.find(qn('p:cNvCxnSpPr'))
                    links = [ends.find(qn(tag)) for tag in ('a:stCxn', 'a:endCxn')]
                    assert all(link is not None for link in links), 'Detached connector'
                    points = [site(shapes[int(link.get('id'))], int(link.get('idx'))) for link in links]
                    (x1, y1), (x2, y2) = points
                    route = [points[0]]
                    if abs(x1-x2) > 0.5 and abs(y1-y2) > 0.5:
                        if int(links[0].get('idx')) in (0, 2):
                            route += [(x1, (y1+y2)/2), (x2, (y1+y2)/2)]
                        else:
                            route += [((x1+x2)/2, y1), ((x1+x2)/2, y2)]
                    route.append(points[1])
                    canvas.setStrokeColor(HexColor('#6b6258'))
                    canvas.setFillColor(HexColor('#6b6258'))
                    canvas.setLineWidth(1)
                    path = canvas.beginPath()
                    path.moveTo(route[0][0], 540-route[0][1])
                    for x, y in route[1:]:
                        path.lineTo(x, 540-y)
                    canvas.drawPath(path)
                    line = shape._element.spPr.find(qn('a:ln'))
                    for tag, tip, prior in (('a:tailEnd', route[-1], route[-2]),
                                            ('a:headEnd', route[0], route[1])):
                        if line.find(qn(tag)) is None:
                            continue
                        dx, dy = tip[0]-prior[0], tip[1]-prior[1]
                        norm = (dx*dx+dy*dy)**0.5
                        dx, dy = dx/norm, dy/norm
                        path = canvas.beginPath()
                        path.moveTo(tip[0], 540-tip[1])
                        for sign in (-1, 1):
                            path.lineTo(tip[0]-5*dx+sign*2*dy, 540-(tip[1]-5*dy-sign*2*dx))
                        path.close()
                        canvas.drawPath(path, fill=1, stroke=0)
                    continue
                x, y, w, h = (value / PX / 2 for value in
                              (shape.left, shape.top, shape.width, shape.height))
                if shape.fill.type == 1:
                    canvas.setFillColor(HexColor('#'+str(shape.fill.fore_color.rgb)))
                    # Read the outline without touching line.color: python-pptx's color getter turns
                    # a no-outline box into a solid outline in the saved deck.
                    dashed = shape.line.fill.type == 1
                    canvas.setStrokeColor(HexColor('#cfc8bd'))
                    canvas.setLineWidth(1)
                    canvas.setDash([6, 4] if dashed else [])
                    canvas.roundRect(x, 540-y-h, w, h, 11, fill=1, stroke=int(dashed))
                    canvas.setDash([])
                if not shape.has_text_frame:
                    continue
                frame = shape.text_frame
                tx = x + frame.margin_left / PX / 2
                ty = y + frame.margin_top / PX / 2
                tw = w - (frame.margin_left+frame.margin_right) / PX / 2
                for paragraph in frame.paragraphs:
                    if not paragraph.runs:
                        continue
                    font = paragraph.runs[0].font
                    name = 'NunitoBold' if font.bold else 'Nunito'
                    size = font.size.pt
                    canvas.setFont(name, size)
                    canvas.setFillColor(HexColor('#'+str(font.color.rgb)))
                    lines = simpleSplit(paragraph.text, name, size, tw)
                    for line in lines:
                        ty += size*1.15
                        assert ty <= y+h+0.1, f'Text overflow: {shape.name}'
                        if paragraph.alignment == 2:
                            canvas.drawCentredString(tx+tw/2, 540-ty, line)
                        else:
                            canvas.drawString(tx, 540-ty, line)
                    ty += paragraph.space_after.pt if paragraph.space_after else 0
            canvas.showPage()
        canvas.save()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--bookends', type=Path, default=ASSETS / 'aura-health-architecture-reviewed.pdf')
    args = parser.parse_args()
    reference = pymupdf.open(args.bookends)
    deck_path = ASSETS / 'aura-health-deck.pptx'
    deck = Presentation(deck_path)
    slides = [deck.slides[number] for number in (10, 11, 12)]
    for slide, kind in zip(slides, KINDS):
        for shape in list(slide.shapes):
            slide.shapes._spTree.remove(shape._element)
        DIAGRAMS[kind](slide)
        marker = '\nArchitecture audit 2026-10-14:'
        notes = slide.notes_slide.notes_text_frame.text.split(marker)[0]
        slide.notes_slide.notes_text_frame.text = notes + (marker + ' implementation 90e31ea. '
            'See docs/submission/ARCHITECTURE_REVIEW.md. The 90-s budget is checked between calls; '
            'resource rendering has no model, but patient-input assessment may use one.')
    with TemporaryDirectory() as temp:
        diagrams_path = Path(temp) / 'diagrams.pdf'
        export_diagrams(slides, diagrams_path)
        diagrams = pymupdf.open(diagrams_path)
        result = pymupdf.open()
        result.insert_pdf(reference, from_page=0, to_page=0)
        result.insert_pdf(diagrams)
        result.insert_pdf(reference, from_page=len(reference)-1, to_page=len(reference)-1)
        reference.close()
        result.set_metadata({'title': 'AURA Health — reviewed architecture',
                             'subject': 'Implementation audit: 90e31ea, 2026-10-14'})
        result.save(ASSETS / 'aura-health-architecture-reviewed.pdf', garbage=4, deflate=True)
        for index, (page, name) in enumerate(zip(diagrams, IMAGES), 11):
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2))
            pixmap.save(ROOT / f'docs/presentation/diagrams/{name}.png')
            pixmap.save(ROOT / f'docs/presentation/deck/png/slide-{index:02d}.png')
    deck.save(deck_path)
    print('Updated 3 native diagram slides, 6 PNGs, the 18-slide deck and 5-page PDF.')


if __name__ == '__main__':
    main()
